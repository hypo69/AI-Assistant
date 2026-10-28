# -*- coding: utf-8 -*-
# =============================================================================
# Process Name: AI-Windows-Admin Apps AI_Assistant - Test Linked Sources & Converters
# =============================================================================
# Description:
#   Тестовый набор для проверки связывания внешних файлов и папок ПК,
#   чтения таблиц Excel (.xlsx), DOCX, CSV, аудиофайлов и Google Sheets,
#   а также их индексации и сквозного поиска знаний БЕЗ перемещения файлов.
#
# File: test_linked_sources_and_converters.py
# Package: apps.ai_assistant.tests
# Author: hypo69
# Copyright: © 2026 hypo69
# Updated: 2026-10-22 17:30:00
# =============================================================================

from pathlib import Path
import pytest
from fastapi.testclient import TestClient
import openpyxl

from apps.ai_assistant.core.content_extractor import ContentExtractor
from apps.ai_assistant.core.knowledge_store import KnowledgeStore
from apps.ai_assistant.core.nlp_analyzers import AnalyzerRegistry
from apps.ai_assistant.core.qa_engine import KnowledgeQAEngine
from apps.ai_assistant.core.temporal_code_graph import TemporalCodeGraph
from apps.ai_assistant.core.workspace_manager import WorkspaceManager
from apps.ai_assistant.main import app
from src.utils.convertors.excel import excel2dict, excel2text
from src.utils.convertors.google_sheet import extract_spreadsheet_id


@pytest.fixture()
def isolated_assistant_env(tmp_path: Path):
    """Инициализирует изолированную среду с SQLite базой, анализаторами и WorkspaceManager."""
    base_dir = tmp_path / "assistant_data"
    registry = AnalyzerRegistry()
    store = KnowledgeStore(db_path=base_dir / "test_store.db")
    graph = TemporalCodeGraph(db_path=base_dir / "test_graph.db")
    extractor = ContentExtractor()
    manager = WorkspaceManager(
        base_dir=base_dir,
        workspaces_dir=base_dir / "workspaces",
        knowledge_store=store,
        analyzer_registry=registry,
        temporal_graph=graph,
        content_extractor=extractor,
    )
    qa_engine = KnowledgeQAEngine(workspace_manager=manager)
    return {
        "tmp_path": tmp_path,
        "base_dir": base_dir,
        "store": store,
        "manager": manager,
        "qa_engine": qa_engine,
        "extractor": extractor,
    }


def test_excel_converter_extraction(tmp_path: Path):
    """Проверяет корректность конвертации таблиц Excel (.xlsx) в словарь и Markdown."""
    excel_file = tmp_path / "sample_books.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Библиотека"
    ws.append(["Автор", "Произведение", "Год"])
    ws.append(["Лев Толстой", "Война и мир", 1869])
    ws.append(["Фёдор Достоевский", "Преступление и наказание", 1866])
    wb.save(str(excel_file))
    wb.close()

    # 1. Проверка excel2dict
    dict_res = excel2dict(excel_file)
    assert "Библиотека" in dict_res
    assert len(dict_res["Библиотека"]) == 2
    assert dict_res["Библиотека"][0]["Автор"] == "Лев Толстой"

    # 2. Проверка excel2text (Markdown)
    md_res = excel2text(excel_file)
    assert "## Лист: Библиотека" in md_res
    assert "| Автор | Произведение | Год |" in md_res
    assert "Война и мир" in md_res


@pytest.mark.asyncio
async def test_docx_and_csv_extraction(tmp_path: Path):
    """Проверяет извлечение структурированного контента из Word (.docx) и CSV."""
    extractor = ContentExtractor()

    # 1. Тест CSV
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("Имя,Статус\nАнна,Активен\nБорис,В ожидании", encoding="utf-8")
    csv_extracted = await extractor.extract(csv_file)
    assert csv_extracted["source_type"] == "csv"
    assert "| Имя | Статус |" in csv_extracted["text"]
    assert "Анна" in csv_extracted["text"]

    # 2. Тест DOCX
    import docx
    docx_file = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("Компания разрабатывает интеллектуальные системы.")
    doc.save(str(docx_file))

    docx_extracted = await extractor.extract(docx_file)
    assert docx_extracted["source_type"] == "docx"
    assert "интеллектуальные системы" in docx_extracted["text"]


def test_google_sheets_url_extraction():
    """Проверяет распознавание и извлечение Spreadsheet ID из ссылок Google Таблиц."""
    url1 = "https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=0"
    assert extract_spreadsheet_id(url1) == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"

    assert ContentExtractor.is_google_sheet_target(url1) is True
    assert ContentExtractor.is_google_sheet_target("C:\\files\\report.pdf") is False


@pytest.mark.asyncio
async def test_workspace_link_sources_lifecycle(isolated_assistant_env):
    """Проверяет жизненный цикл связывания источников: добавление, список и удаление."""
    manager: WorkspaceManager = isolated_assistant_env["manager"]
    tmp_path: Path = isolated_assistant_env["tmp_path"]

    # Создаём внешние файлы и папки
    external_file = tmp_path / "external_manual.txt"
    external_file.write_text("Инструкция пользователя.", encoding="utf-8")

    external_dir = tmp_path / "external_folder"
    external_dir.mkdir(parents=True, exist_ok=True)

    ws_id = "test_ws"

    # Привязка файла
    src_file = manager.link_source(ws_id, str(external_file), name="Руководство")
    assert src_file["source_type"] == "file"
    assert src_file["source_path"] == str(external_file)

    # Привязка директории
    src_dir = manager.link_source(ws_id, str(external_dir), name="Внешняя папка")
    assert src_dir["source_type"] == "directory"

    # Привязка ссылки Google Таблицы
    sheet_url = "https://docs.google.com/spreadsheets/d/1234567890abcdefghij/edit"
    src_sheet = manager.link_source(ws_id, sheet_url, name="Сводная таблица")
    assert src_sheet["source_type"] == "google_sheet"

    # Список источников
    sources = manager.list_sources(ws_id)
    assert len(sources) == 3

    # Удаление источника
    deleted = manager.unlink_source(ws_id, src_file["source_id"])
    assert deleted is True
    assert len(manager.list_sources(ws_id)) == 2


@pytest.mark.asyncio
async def test_sync_linked_sources_without_moving_files(isolated_assistant_env):
    """Проверяет, что файлы сканируются и индексируются БЕЗ перемещения и копирования оригинала."""
    manager: WorkspaceManager = isolated_assistant_env["manager"]
    qa_engine: KnowledgeQAEngine = isolated_assistant_env["qa_engine"]
    tmp_path: Path = isolated_assistant_env["tmp_path"]

    # Создаём внешний каталог, имитирующий пользовательские файлы на диске
    external_dir = tmp_path / "user_documents_on_c_drive"
    external_dir.mkdir(parents=True, exist_ok=True)

    file_a = external_dir / "mama_linked.txt"
    file_a.write_text("Мама мыла раму.", encoding="utf-8")

    file_b = external_dir / "farm_linked.txt"
    file_b.write_text("לדוד משה היתה חווה", encoding="utf-8")

    ws_id = "ws_synced"
    ws = manager.create_or_select_workspace(ws_id, name="Связанная область")
    ws_path = Path(ws["path"])

    # Привязываем внешнюю папку
    manager.link_source(ws_id, str(external_dir), name="Мои документы")

    # Запускаем синхронизацию
    sync_report = await manager.sync_linked_sources(ws_id)
    assert sync_report["status"] == "ok"
    assert sync_report["total_indexed"] == 2

    # КРИТИЧЕСКАЯ ПРОВЕРКА: файлы НЕ перемещены и НЕ скопированы в папку репозитория workspace!
    assert not (ws_path / "mama_linked.txt").exists(), "Файл не должен копироваться в папку workspace!"
    assert not (ws_path / "farm_linked.txt").exists(), "Файл не должен копироваться в папку workspace!"
    assert file_a.exists(), "Оригинальный файл должен оставаться на месте"
    assert file_b.exists(), "Оригинальный файл должен оставаться на месте"

    # Проверка, что документы зарегистрированы в SQLite с указанием оригинального пути
    docs = manager.store.list_documents(ws_id)
    assert len(docs) == 2
    doc_paths = [d["file_path"] for d in docs]
    assert str(file_a) in doc_paths
    assert str(file_b) in doc_paths

    # Проверка, что извлечены факты
    facts = manager.store.get_facts(workspace_id=ws_id)
    assert len(facts) >= 2

    # Проверка, что QA-движок отвечает на вопросы по этим связанным файлам
    ans_ru = await qa_engine.ask(
        question="Что мыла мама?",
        workspace_id=ws_id,
        mode="local_only",
    )
    assert "раму" in ans_ru["answer"].lower()

    ans_he = await qa_engine.ask(
        question="למי הייתה חווה?",
        workspace_id=ws_id,
        mode="local_only",
    )
    assert "לדוד משה" in ans_he["answer"]


def test_rest_api_sources_endpoints(isolated_assistant_env):
    """Интеграционный тест REST API endpoints управления связанными источниками."""
    client = TestClient(app)
    tmp_path: Path = isolated_assistant_env["tmp_path"]

    test_doc = tmp_path / "api_test_doc.txt"
    test_doc.write_text("Тестовый документ для API.", encoding="utf-8")

    ws_id = "api_ws"

    # 1. Привязка источника через API
    res_link = client.post(
        f"/api/v1/workspaces/{ws_id}/sources",
        json={"target": str(test_doc), "name": "API Тест", "auto_watch": True},
    )
    assert res_link.status_code == 200
    data_link = res_link.json()
    assert data_link["status"] == "ok"
    src_id = data_link["source"]["source_id"]

    # 2. Получение списка источников
    res_list = client.get(f"/api/v1/workspaces/{ws_id}/sources")
    assert res_list.status_code == 200
    assert len(res_list.json()["sources"]) >= 1

    # 3. Синхронизация источников через API
    res_sync = client.post(f"/api/v1/workspaces/{ws_id}/sources/sync")
    assert res_sync.status_code == 200
    assert res_sync.json()["status"] == "ok"

    # 4. Предпросмотр извлечения контента
    res_prev = client.post(
        "/api/v1/extract-preview",
        json={"target": str(test_doc)},
    )
    assert res_prev.status_code == 200
    assert "Тестовый документ" in res_prev.json()["text"]

    # 5. Удаление источника
    res_del = client.delete(f"/api/v1/workspaces/{ws_id}/sources/{src_id}")
    assert res_del.status_code == 200
    assert res_del.json()["deleted"] is True
