# -*- coding: utf-8 -*-
# =============================================================================
# Process Name: AI-Windows-Admin Apps AI_Assistant Core - Content Extractor
# =============================================================================
# Description:
#   Универсальный конвейер извлечения и конвертации разнородных типов данных
#   (PDF, Изображения/OCR, Аудио, Excel, CSV, DOCX, Google Sheets, Markdown, Код)
#   в структурированный текст. Поддерживает переключение между штатными
#   локальными конвертерами и расшифровкой мультимодальными моделями (Gemini).
#
# Usage Examples:
#   Python API:
#     from apps.ai_assistant.core.content_extractor import ContentExtractor
#
#     extractor = ContentExtractor()
#     res = await extractor.extract("docs/sample.pdf", pdf_mode="model")
#     print(res["text"])
#
# File: content_extractor.py
# Project: ai-windows-admin
# Package: apps.ai_assistant.core
# Author: hypo69
# Copyright: © 2026 hypo69
# Updated: 2026-10-26 00:05:00
# =============================================================================

from __future__ import annotations
"""Универсальный экстрактор и конвертер разнородных документов, изображений и медиа в текст."""

import asyncio
import csv
import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from logger import logger
from src.utils.convertors.audio import _get_gemini_client, audio2text
from src.utils.convertors.excel import excel2text
from src.utils.convertors.google_sheet import extract_spreadsheet_id, read_google_sheet
from src.utils.pdf import extract_pdf_text

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tiff", ".gif"}
AUDIO_EXTENSIONS = {".wav", ".mp3", ".ogg", ".m4a", ".flac", ".aac"}
PDF_EXTENSIONS = {".pdf"}
TABLE_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls"}
DOC_EXTENSIONS = {".docx"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".log", ".yaml", ".yml", ".py", ".html", ".htm",
    ".js", ".ts", ".css", ".sql", ".ps1", ".bat", ".sh", ".xml", ".ini", ".toml", ".cfg", ".rst",
}

SUPPORTED_EXTENSIONS = (
    IMAGE_EXTENSIONS
    | AUDIO_EXTENSIONS
    | PDF_EXTENSIONS
    | TABLE_EXTENSIONS
    | DOC_EXTENSIONS
    | TEXT_EXTENSIONS
)

IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".tiff": "image/tiff",
}


class ContentExtractor:
    """Универсальный сервис конвертации разноформатных данных в текст с поддержкой локального и ИИ-режимов."""

    def __init__(self) -> None:
        """Инициализирует конвертер контента."""
        pass

    @staticmethod
    def is_google_sheet_target(target: str) -> bool:
        """Проверяет, является ли цель ссылкой на Google Таблицу."""
        t = target.strip().lower()
        return "docs.google.com/spreadsheets" in t or "google.com/spreadsheets" in t

    @staticmethod
    def is_supported_file(path_or_name: Union[str, Path]) -> bool:
        """Проверяет, поддерживается ли расширение файла."""
        ext = Path(path_or_name).suffix.lower()
        return ext in SUPPORTED_EXTENSIONS

    def _extract_docx(self, file_path: Path) -> str:
        """Извлекает текст и таблицы из документа Microsoft Word (.docx)."""
        try:
            import docx

            doc = docx.Document(str(file_path))
            elements: List[str] = [f"# Документ Word: {file_path.name}\n"]

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    elements.append(p_text)

            for table_idx, table in enumerate(doc.tables):
                elements.append(f"\n### Таблица {table_idx + 1}")
                table_rows: List[List[str]] = []
                for row in table.rows:
                    row_cells = [c.text.strip().replace("\n", " ").replace("|", "/") for c in row.cells]
                    table_rows.append(row_cells)

                if table_rows:
                    headers = table_rows[0]
                    elements.append("| " + " | ".join(headers) + " |")
                    elements.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    for row in table_rows[1:]:
                        elements.append("| " + " | ".join(row) + " |")
                    elements.append("")

            return "\n\n".join(elements)
        except Exception as exc:
            logger.error(f"[ContentExtractor] Ошибка разбора DOCX {file_path}: {exc}")
            return f"*(Ошибка чтения документа Word {file_path.name}: {exc})*"

    def _extract_csv(self, file_path: Path) -> str:
        """Извлекает и форматирует файл CSV в Markdown-таблицу."""
        try:
            text = file_path.read_text(encoding="utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text))
            rows = [r for r in reader if any(cell.strip() for cell in r)]
            if not rows:
                return f"# Файл CSV: {file_path.name}\n\n*(Файл пуст)*"

            headers = [c if c.strip() else f"Колонка_{idx + 1}" for idx, c in enumerate(rows[0])]
            lines = [
                f"# Таблица CSV: {file_path.name}",
                "",
                "| " + " | ".join(headers) + " |",
                "| " + " | ".join(["---"] * len(headers)) + " |",
            ]
            for row in rows[1:]:
                padded = row + [""] * (len(headers) - len(row))
                row_vals = [c.replace("\n", " ").replace("|", "/") for c in padded[:len(headers)]]
                lines.append("| " + " | ".join(row_vals) + " |")

            return "\n".join(lines)
        except Exception as exc:
            logger.error(f"[ContentExtractor] Ошибка разбора CSV {file_path}: {exc}")
            return f"*(Ошибка чтения CSV {file_path.name}: {exc})*"

    def _extract_plain_text(self, file_path: Path) -> str:
        """Читает обычный текстовый файл с авто-определением кодировки (UTF-8, CP1251, Latin-1)."""
        encodings = ["utf-8", "cp1251", "windows-1252", "latin-1"]
        raw_bytes = file_path.read_bytes()
        for enc in encodings:
            try:
                return raw_bytes.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("utf-8", errors="replace")

    async def convert_pdf(self, file_path: Path, mode: str = "local") -> str:
        """Конвертирует PDF в Markdown локальными средствами или через мультимодальную модель Gemini.

        Args:
            file_path: Путь к файлу PDF.
            mode: 'local' (pypdf/pdfminer) или 'model' (Gemini).
        """
        if mode == "model":
            client = _get_gemini_client()
            if client:
                try:
                    from google.genai import types

                    pdf_bytes = file_path.read_bytes()
                    prompt = (
                        "Преобразуй содержимое этого PDF-документа в аккуратный и структурированный Markdown на русском языке. "
                        "Сохраняй структуру разделов, списки, формулы и таблицы. "
                        "Выведи только распознанный Markdown-текст без вводных фраз."
                    )
                    part = types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf")
                    resp = await asyncio.to_thread(
                        client.models.generate_content,
                        model="gemini-2.5-flash",
                        contents=[part, prompt],
                    )
                    if resp and resp.text:
                        return resp.text.strip()
                except Exception as exc:
                    logger.warning(f"[ContentExtractor] Сбой модели Gemini для PDF {file_path.name}, fallback на локальный: {exc}")

        # Штатный локальный экстрактор (pypdf -> pdfminer)
        pdf_text = extract_pdf_text(file_path)
        return pdf_text or f"*(Пустой или нераспознанный PDF-документ: {file_path.name})*"

    async def convert_image_ocr(self, file_path: Path, mode: str = "model") -> str:
        """Выполняет оптическое распознавание символов (OCR) для изображений (PNG, JPG, BMP и др.).

        Args:
            file_path: Путь к файлу изображения.
            mode: 'model' (Gemini Vision) или 'local' (Tesseract / PIL).
        """
        if mode == "model":
            client = _get_gemini_client()
            if client:
                try:
                    from google.genai import types

                    img_bytes = file_path.read_bytes()
                    mime = IMAGE_MIME_TYPES.get(file_path.suffix.lower(), "image/png")
                    prompt = (
                        "Выполни оптическое распознавание текста (OCR) этого изображения. "
                        "Извлеки весь видимый текст, таблицы, надписи и структурируй в аккуратном формате Markdown на русском языке. "
                        "Выведи только распознанный текст без вступительных или служебных фраз."
                    )
                    part = types.Part.from_bytes(data=img_bytes, mime_type=mime)
                    resp = await asyncio.to_thread(
                        client.models.generate_content,
                        model="gemini-2.5-flash",
                        contents=[part, prompt],
                    )
                    if resp and resp.text:
                        return resp.text.strip()
                except Exception as exc:
                    logger.warning(f"[ContentExtractor] Сбой модели Gemini Vision для OCR {file_path.name}: {exc}")

        # Локальное штатное распознавание: попытка вызвать Tesseract OCR
        try:
            import pytesseract
            from PIL import Image

            img = Image.open(file_path)
            ocr_res = pytesseract.image_to_string(img, lang="rus+eng").strip()
            if ocr_res:
                return f"# OCR (Tesseract): {file_path.name}\n\n{ocr_res}"
        except Exception:
            pass

        # Если локального Tesseract нет — формируем карточку метаданных изображения через PIL
        try:
            from PIL import Image

            with Image.open(file_path) as img:
                w, h = img.size
                fmt = img.format or file_path.suffix.upper().replace(".", "")
                color_mode = img.mode
            return (
                f"# Изображение: {file_path.name}\n\n"
                f"- **Формат:** {fmt}\n"
                f"- **Разрешение:** {w}x{h} px\n"
                f"- **Цветовая модель:** {color_mode}\n\n"
                f"*(Локальный модуль Tesseract OCR не установлен в системе. "
                f"Для автоматического извлечения текста выберите режим 'Расшифровка на модели (Gemini)')*"
            )
        except Exception as exc:
            return f"# Изображение: {file_path.name}\n\n*(Ошибка чтения изображения: {exc})*"

    async def convert_audio(self, file_path: Path, mode: str = "model", language: str = "ru") -> str:
        """Выполняет транскрибацию аудиофайла строго с подключением к внешней модели Gemini (или пропуск при raw).

        Args:
            file_path: Путь к аудиофайлу.
            mode: Режим распознавания ('model' или 'raw').
            language: Язык аудио.
        """
        if mode in ("raw", "none"):
            return ""
        # По ТЗ: транскрибацию аудио делать только с подключением к внешней модели
        logger.info(f"[ContentExtractor] Транскрибация аудио через модель Gemini: {file_path.name}")
        audio_text = await audio2text(file_path, language=language)
        return audio_text or f"*(Не удалось транскрибировать аудиозапись: {file_path.name})*"

    async def extract(
        self,
        target: Union[str, Path],
        language: str = "ru",
        pdf_mode: str = "local",
        ocr_mode: str = "model",
        audio_mode: str = "model",
    ) -> Dict[str, Any]:
        """Универсально извлекает текстовое содержимое из локального файла или Google Sheets.

        Args:
            target: Путь к файлу на компьютере или ссылка на Google Sheets.
            language: Предпочтительный язык для аудио транскрибации.
            pdf_mode: Режим конвертации PDF ('local' или 'model').
            ocr_mode: Режим оптического распознавания изображений ('local' или 'model').
            audio_mode: Режим транскрибации аудио ('model').

        Returns:
            Dict[str, Any]: Словарь с извлечённым текстом, типом источника и метаданными.
        """
        target_str = str(target).strip()

        # 1. Проверка на Google Таблицу
        if self.is_google_sheet_target(target_str):
            logger.info(f"[ContentExtractor] Чтение Google Таблицы: {target_str}")
            sheet_text = await read_google_sheet(target_str)
            sid = extract_spreadsheet_id(target_str)
            return {
                "text": sheet_text,
                "source_type": "google_sheet",
                "source_path": target_str,
                "document_id": f"gsheet_{sid}",
                "title": f"Google Sheet {sid}",
                "extension": ".gsheet",
            }

        file_path = Path(target).resolve()
        if not file_path.exists():
            raise FileNotFoundError(f"Файл не найден: {file_path}")

        suffix = file_path.suffix.lower()

        # 2. PDF файл (локальный или расшифровка на модели)
        if suffix in PDF_EXTENSIONS:
            if pdf_mode in ("raw", "none"):
                return {
                    "text": "",
                    "source_type": "pdf",
                    "source_path": str(file_path),
                    "document_id": file_path.name,
                    "title": file_path.stem,
                    "extension": suffix,
                    "conversion_mode": "raw",
                    "is_raw": True,
                }
            logger.info(f"[ContentExtractor] Конвертация PDF ({pdf_mode}): {file_path.name}")
            pdf_text = await self.convert_pdf(file_path, mode=pdf_mode)
            return {
                "text": pdf_text,
                "source_type": "pdf",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": file_path.stem,
                "extension": suffix,
                "conversion_mode": pdf_mode,
                "is_raw": False,
            }

        # 3. Изображения (OCR локальный или на модели)
        if suffix in IMAGE_EXTENSIONS:
            if ocr_mode in ("raw", "none"):
                return {
                    "text": "",
                    "source_type": "image_ocr",
                    "source_path": str(file_path),
                    "document_id": file_path.name,
                    "title": f"Изображение: {file_path.name}",
                    "extension": suffix,
                    "conversion_mode": "raw",
                    "is_raw": True,
                }
            logger.info(f"[ContentExtractor] OCR изображения ({ocr_mode}): {file_path.name}")
            ocr_text = await self.convert_image_ocr(file_path, mode=ocr_mode)
            return {
                "text": ocr_text,
                "source_type": "image_ocr",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": f"OCR: {file_path.name}",
                "extension": suffix,
                "conversion_mode": ocr_mode,
                "is_raw": False,
            }

        # 4. Аудиофайлы (транскрибация внешней моделью Gemini или сохранение в сыром виде)
        if suffix in AUDIO_EXTENSIONS:
            if audio_mode in ("raw", "none"):
                return {
                    "text": "",
                    "source_type": "audio",
                    "source_path": str(file_path),
                    "document_id": file_path.name,
                    "title": f"Аудио (сырой): {file_path.stem}",
                    "extension": suffix,
                    "conversion_mode": "raw",
                    "is_raw": True,
                }
            logger.info(f"[ContentExtractor] Транскрибация аудио ({audio_mode}): {file_path.name}")
            audio_text = await self.convert_audio(file_path, mode=audio_mode, language=language)
            return {
                "text": audio_text,
                "source_type": "audio",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": f"Транскрипция аудио: {file_path.stem}",
                "extension": suffix,
                "conversion_mode": audio_mode,
                "is_raw": False,
            }

        # 5. Excel файл (.xlsx, .xls)
        if suffix in (".xlsx", ".xls"):
            logger.info(f"[ContentExtractor] Конвертация Excel в Markdown: {file_path.name}")
            excel_text = excel2text(file_path)
            return {
                "text": excel_text,
                "source_type": "excel",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": file_path.stem,
                "extension": suffix,
            }

        # 6. Microsoft Word (.docx)
        if suffix in DOC_EXTENSIONS:
            logger.info(f"[ContentExtractor] Извлечение Word DOCX: {file_path.name}")
            docx_text = self._extract_docx(file_path)
            return {
                "text": docx_text,
                "source_type": "docx",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": file_path.stem,
                "extension": suffix,
            }

        # 7. CSV / TSV
        if suffix in (".csv", ".tsv"):
            csv_text = self._extract_csv(file_path)
            return {
                "text": csv_text,
                "source_type": "csv",
                "source_path": str(file_path),
                "document_id": file_path.name,
                "title": file_path.stem,
                "extension": suffix,
            }

        # 8. Обычный текст / Исходный код / JSON / Markdown
        plain_text = self._extract_plain_text(file_path)
        return {
            "text": plain_text,
            "source_type": "text",
            "source_path": str(file_path),
            "document_id": file_path.name,
            "title": file_path.stem,
            "extension": suffix,
        }
