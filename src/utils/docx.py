# -*- coding: utf-8 -*-
# =============================================================================
# Process Name: AI Assistant UTILS - DOCX Text Extraction
# =============================================================================
# Description:
#   Извлечение текстового содержимого и форматированного markdown из файлов .docx.
#
# File: docx.py
# Package: src.utils
# Author: hypo69
# Copyright: © 2026 hypo69
# Updated: 2026-10-28 18:26:00
# =============================================================================

from __future__ import annotations

"""Модуль безопасного извлечения текста из документов Microsoft Word (.docx)."""

from pathlib import Path
from typing import Optional, Union

from logger import logger

try:
    import docx
except ImportError:
    docx = None


def extract_docx_text(file_path: Union[str, Path], as_markdown: bool = False) -> Optional[str]:
    """Извлекает текст из документа Word (.docx).

    Args:
        file_path (Union[str, Path]): Путь к файлу .docx.
        as_markdown (bool): Флаг форматирования заголовков и списков в стиле Markdown.

    Returns:
        Optional[str]: Извлечённый текст или None в случае ошибки или отсутствия файла.
    """
    path = Path(file_path)
    if not path.is_file():
        logger.warning(f"[extract_docx_text] Файл не найден: {path}")
        return None

    if docx is None:
        logger.warning("[extract_docx_text] Библиотека python-docx не установлена.")
        return None

    try:
        doc = docx.Document(str(path))
        lines: list[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = (p.style.name or "").lower() if p.style else ""
            if as_markdown:
                if "heading 1" in style_name:
                    lines.append(f"# {text}")
                elif "heading 2" in style_name:
                    lines.append(f"## {text}")
                elif "heading 3" in style_name:
                    lines.append(f"### {text}")
                elif "list" in style_name or "bullet" in style_name:
                    lines.append(f"- {text}")
                else:
                    lines.append(text)
            else:
                lines.append(text)

        # Обработка таблиц
        for table in doc.tables:
            table_rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    if as_markdown:
                        table_rows.append("| " + " | ".join(cells) + " |")
                    else:
                        table_rows.append("\t".join(cells))
            if table_rows:
                lines.append("\n".join(table_rows))

        return "\n\n".join(lines)
    except Exception as exc:
        logger.error(f"[extract_docx_text] Ошибка извлечения текста из {path.name}: {exc}")
        return None
